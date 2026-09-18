"""Source ingestion for /learn.

Keeps file/folder parsing separate from the existing parameter-update path.
Text extraction is deterministic; media files are only classified here and are not
pretended to be understood by a text-only controller.
"""
from __future__ import annotations

import csv
import json
import os
from pathlib import Path
from typing import Any, Dict, Iterable, List

TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".toml",
    ".yaml", ".yml", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cc", ".cpp", ".h",
    ".hpp", ".rs", ".go", ".swift", ".kt", ".kts", ".sh", ".zsh", ".ps1",
    ".sql", ".css", ".scss",
}
STRUCTURED_SUFFIXES = {".json", ".jsonl", ".csv", ".tsv", ".parquet"}
DOCUMENT_SUFFIXES = {".pdf", ".docx"}
MEDIA_SUFFIXES = {
    "image": {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tiff", ".tif"},
    "video": {".mp4", ".mov", ".mkv", ".webm", ".avi", ".m4v"},
    "audio": {".wav", ".mp3", ".flac", ".m4a", ".ogg", ".aac", ".opus"},
}
SUPPORTED_TEXT_SUFFIXES = TEXT_SUFFIXES | STRUCTURED_SUFFIXES | DOCUMENT_SUFFIXES


def media_kind(path: Path) -> str | None:
    suffix = path.suffix.lower()
    for kind, suffixes in MEDIA_SUFFIXES.items():
        if suffix in suffixes:
            return kind
    return None


def _iter_source_files(source: Path, max_files: int = 128) -> Iterable[Path]:
    if source.is_file():
        yield source
        return
    count = 0
    for path in sorted(source.rglob("*")):
        if path.is_file():
            yield path
            count += 1
            if count >= max_files:
                break


def inspect_learning_source(source: str | os.PathLike[str]) -> Dict[str, Any]:
    path = Path(source).expanduser().resolve()
    if not path.exists():
        raise ValueError(f"Learning source does not exist: {path}")
    text_files: List[str] = []
    media_files: Dict[str, List[str]] = {"image": [], "video": [], "audio": []}
    ignored: List[str] = []
    for file in _iter_source_files(path):
        suffix = file.suffix.lower()
        kind = media_kind(file)
        if kind:
            media_files[kind].append(str(file))
        elif suffix in SUPPORTED_TEXT_SUFFIXES:
            text_files.append(str(file))
        else:
            ignored.append(str(file))
    return {
        "path": str(path),
        "is_folder": path.is_dir(),
        "text_files": text_files,
        "media_files": media_files,
        "media_kinds": [kind for kind, files in media_files.items() if files],
        "ignored_files": ignored,
    }


def _clip(text: str, max_chars: int) -> str:
    text = str(text or "").strip()
    return text if len(text) <= max_chars else text[:max_chars] + "\n[...truncated...]"


def _extract_json(path: Path) -> str:
    obj = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    return json.dumps(obj, ensure_ascii=False, indent=2)


def _extract_jsonl(path: Path) -> str:
    rows = []
    for index, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines()):
        if not line.strip():
            continue
        try:
            rows.append(json.loads(line))
        except Exception:
            rows.append({"raw": line})
        if index >= 1999:
            break
    return "\n".join(json.dumps(row, ensure_ascii=False) for row in rows)


def _extract_delimited(path: Path, delimiter: str) -> str:
    out: List[str] = []
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for index, row in enumerate(reader):
            out.append(delimiter.join(str(cell) for cell in row))
            if index >= 1999:
                break
    return "\n".join(out)


def _extract_parquet(path: Path) -> str:
    try:
        import pyarrow.parquet as pq
    except Exception as exc:
        raise RuntimeError("Parquet /learn requires pyarrow") from exc
    table = pq.read_table(path)
    if table.num_rows > 2000:
        table = table.slice(0, 2000)
    rows = table.to_pylist()
    return "\n".join(json.dumps(row, ensure_ascii=False, default=str) for row in rows)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise RuntimeError("PDF /learn requires pypdf") from exc
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages[:200]:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("DOCX /learn requires python-docx") from exc
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text_file(path: Path, max_chars: int = 120_000) -> str:
    suffix = path.suffix.lower()
    if suffix == ".json":
        text = _extract_json(path)
    elif suffix == ".jsonl":
        text = _extract_jsonl(path)
    elif suffix == ".csv":
        text = _extract_delimited(path, ",")
    elif suffix == ".tsv":
        text = _extract_delimited(path, "\t")
    elif suffix == ".parquet":
        text = _extract_parquet(path)
    elif suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix == ".docx":
        text = _extract_docx(path)
    elif suffix in TEXT_SUFFIXES:
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        return ""
    return _clip(text, max_chars)


def extract_learning_text(
    source: str | os.PathLike[str],
    *,
    max_total_chars: int = 240_000,
    max_chars_per_file: int = 120_000,
) -> Dict[str, Any]:
    inspection = inspect_learning_source(source)
    chunks: List[str] = []
    total = 0
    errors: List[str] = []
    for raw in inspection["text_files"]:
        path = Path(raw)
        try:
            text = extract_text_file(path, max_chars=max_chars_per_file)
        except Exception as exc:
            errors.append(f"{path.name}: {type(exc).__name__}: {exc}")
            continue
        if not text.strip():
            continue
        remaining = max_total_chars - total
        if remaining <= 0:
            break
        text = text[:remaining]
        chunks.append(f"===== SOURCE FILE: {path.name} =====\n{text}")
        total += len(text)
    return {
        **inspection,
        "text": "\n\n".join(chunks),
        "text_chars": total,
        "extraction_errors": errors,
    }
